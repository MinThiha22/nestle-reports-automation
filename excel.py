from openpyxl import load_workbook
import datetime
import sys

'''
this is a script to update the A2 cell in an Excel file 
by adding 7 days to the current date in that cell using 
openpyxl library, to be use only in ElectronHumby project.
'''
"""Update A2 cell date by adding 7 days using openpyxl"""
file_path = r"C:\Users\NZShallaZu\NESTLE\NZ CDT - NEW - Documents\General\1. Reporting\Power BI\Store Level\Data Dump\FSCalendar.xlsx"
# file_path = r"C:\Users\chg\Downloads\FSCalendar3.xlsx"
try:
    # Load the workbook
    wb = load_workbook(file_path)
    ws = wb.active
    current_date_in_a2 = ws['A2'].value
    # --- Step 2: Calculate the date for "next week" in Python ---
    next_week_date = current_date_in_a2 + datetime.timedelta(days=7)
    ws['A2'].value = next_week_date
    
    wb.save(file_path)
    sys.exit(0)
    
except Exception as e:
    sys.exit(1)
